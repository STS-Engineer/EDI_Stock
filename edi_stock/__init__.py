"""Application factory. Importing this package never opens a database connection."""
import hmac
import io
import os
import re
import secrets
import uuid
from pathlib import Path

from flask import Flask, Response, abort, jsonify, redirect, render_template, request, session, url_for
from openpyxl import Workbook
from openpyxl.comments import Comment
from sqlalchemy.exc import DataError, IntegrityError, InterfaceError, OperationalError, SQLAlchemyError, TimeoutError as SQLAlchemyTimeoutError
from werkzeug.exceptions import HTTPException

from .parsers import parse_upload
from .repository import ImportConflict, InventoryConflict, PostgresRepository
from .staging import PreviewStore
from .validation import SCHEMAS, ValidationError, validate_rows


def create_app(config=None, *, repository=None):
    app = Flask(__name__)
    app.config.from_mapping(
        SECRET_KEY=os.environ.get('SECRET_KEY'), DATABASE_URL=os.environ.get('DATABASE_URL'),
        IMPORT_API_TOKEN=os.environ.get('IMPORT_API_TOKEN'),
        UI_IMPORTS_ENABLED=os.environ.get('UI_IMPORTS_ENABLED', 'false').lower() == 'true',
        MAX_CONTENT_LENGTH=16 * 1024 * 1024, MAX_FORM_MEMORY_SIZE=100000, MAX_FORM_PARTS=10,
        MAX_IMPORT_ROWS=10000, PREVIEW_TTL=3600,
        PREVIEW_DIR=os.environ.get('PREVIEW_DIR', str(Path(app.instance_path) / 'previews')),
        SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax',
        SESSION_COOKIE_SECURE=os.environ.get('COOKIE_SECURE', 'true').lower() != 'false',
    )
    if config:
        app.config.update(config)
    repo = repository
    if repo is None and app.config['DATABASE_URL']:
        repo = PostgresRepository(app.config['DATABASE_URL'])
    app.extensions['import_repository'] = repo

    def browser_state():
        if not app.secret_key:
            abort(503, description='La clé de session serveur doit être configurée avant de préparer un import.')
        if 'owner' not in session:
            session['owner'] = secrets.token_urlsafe(32)
            session['csrf'] = secrets.token_urlsafe(32)
        return session['owner'], session['csrf']

    def store():
        browser_state()
        return PreviewStore(app.config['PREVIEW_DIR'], app.secret_key, app.config['PREVIEW_TTL'])

    def check_csrf():
        _, expected = browser_state()
        actual = request.form.get('csrf_token', '')
        if not hmac.compare_digest(actual.encode(), expected.encode()):
            abort(403, description='Session de formulaire expirée. Rechargez la page et réessayez.')

    def page(status=200, **values):
        defaults = dict(active_tab='LIVRAISON', csrf_token='', configured=repo is not None and app.config['UI_IMPORTS_ENABLED'],
                        message=None, error=False, errors=[], preview=None, receipt=None)
        if app.secret_key:
            defaults['csrf_token'] = browser_state()[1]
        defaults.update(values)
        return render_template('index.html', **defaults), status

    def error_response(code, message, status, *, retryable=False, errors=None, preview=None):
        if request.path.startswith('/api/'):
            result = jsonify({'status': 'error', 'error': {'code': code, 'message': message,
                              'retryable': retryable, 'details': errors or []}, 'request_id': request.environ['edi.request_id']})
            result.status_code = status
            if retryable:
                result.headers['Retry-After'] = '30'
            return result
        try:
            file_type = request.form.get('file_type', 'LIVRAISON') if request.mimetype != 'application/json' else 'LIVRAISON'
        except HTTPException:
            file_type = 'LIVRAISON'
        return page(status, active_tab=file_type if file_type in SCHEMAS else 'LIVRAISON', message=message, error=True, errors=errors or [], preview=preview)

    @app.before_request
    def request_identity():
        request.environ['edi.request_id'] = uuid.uuid4().hex

    @app.after_request
    def response_headers(response):
        response.headers['X-Request-ID'] = request.environ.get('edi.request_id', '')
        response.headers['X-Content-Type-Options'] = 'nosniff'
        response.headers['Referrer-Policy'] = 'same-origin'
        response.headers['Content-Security-Policy'] = "default-src 'self'; style-src 'self'; script-src 'self'; img-src 'self' data:; object-src 'none'; base-uri 'self'; form-action 'self'; frame-ancestors 'none'"
        if request.path != '/static/app.css' and request.path != '/static/app.js':
            response.headers['Cache-Control'] = 'no-store'
        return response

    @app.errorhandler(ValidationError)
    def invalid(exc):
        return error_response('validation_failed', str(exc), 422, errors=exc.errors)

    @app.errorhandler(ImportConflict)
    def idempotency_conflict(exc):
        return error_response('idempotency_conflict', str(exc), 409)

    @app.errorhandler(InventoryConflict)
    def inventory_conflict(exc):
        return error_response('inventory_conflict', str(exc), 409)

    @app.errorhandler(DataError)
    def database_data_failed(exc):
        return error_response('database_data_error', 'Les valeurs ne correspondent pas au schéma de la base. Vérification requise.', 422)

    @app.errorhandler(IntegrityError)
    def database_constraint_failed(exc):
        return error_response('database_constraint_conflict', 'Une contrainte de la base empêche cet import. Réconciliation requise.', 409)

    @app.errorhandler(SQLAlchemyError)
    def database_failed(exc):
        # SQL exceptions can contain parameter data and credentials. Never serialize/log them.
        app.logger.error('Import database failure type=%s request_id=%s', type(exc).__name__, request.environ['edi.request_id'])
        retryable = isinstance(exc, (OperationalError, InterfaceError, SQLAlchemyTimeoutError))
        if not retryable:
            return error_response('database_schema_error', 'Configuration ou schéma de base incompatible. Intervention opérateur requise.', 503)
        preview_data = None
        if request.path == '/insert':
            try:
                token = request.form.get('temp_file', '')
                _, saved = store().load(token, browser_state()[0], request.form.get('file_type'))
                preview_data = {'file_type': saved['file_type'], 'rows': saved['rows'][:20],
                                'columns': SCHEMAS[saved['file_type']], 'row_count': len(saved['rows']), 'token': token}
            except ValidationError:
                pass
        return error_response('database_unavailable', 'Import non confirmé. Réessayez avec la même clé ou le même aperçu.',
                              503, retryable=True, preview=preview_data)

    @app.errorhandler(HTTPException)
    def http_failed(exc):
        code = {400: 'bad_request', 401: 'unauthorized', 403: 'forbidden', 404: 'not_found',
                413: 'payload_too_large', 415: 'unsupported_media_type', 503: 'not_configured'}.get(exc.code, 'http_error')
        message = 'Le fichier dépasse la limite de 16 Mo.' if exc.code == 413 else exc.description
        return error_response(code, message, exc.code)

    @app.errorhandler(Exception)
    def unexpected(exc):
        if app.testing:
            raise exc
        app.logger.error('Import unexpected failure type=%s request_id=%s', type(exc).__name__, request.environ['edi.request_id'])
        return error_response('internal_error', 'Le traitement a échoué. Conservez le fichier et contactez le support avec la référence de requête.', 500)

    @app.get('/')
    def index():
        tab = request.args.get('tab', 'LIVRAISON')
        return page(active_tab=tab if tab in SCHEMAS else 'LIVRAISON')

    @app.get('/healthz')
    def health():
        return jsonify({'status': 'ok'})

    @app.post('/preview')
    def preview():
        check_csrf()
        file_type = request.form.get('file_type')
        if file_type not in SCHEMAS:
            raise ValidationError('Type de fichier attendu : EDI ou LIVRAISON.')
        upload = request.files.get('file')
        if upload is None or not upload.filename:
            raise ValidationError('Sélectionnez un fichier à importer.')
        rows = parse_upload(upload.read(), upload.filename, file_type, max_rows=app.config['MAX_IMPORT_ROWS'])
        rows = validate_rows(rows, file_type, max_rows=app.config['MAX_IMPORT_ROWS'])
        token = store().save(browser_state()[0], file_type, rows)
        message = f'{len(rows)} lignes validées. Vérifiez les données avant confirmation.'
        if upload.filename.lower().endswith('.pdf'):
            message += ' PDF : comparez chaque page, référence et quantité avec le document original ; une extraction partielle reste possible. Site Tunisia et statut Dispatched sont présélectionnés.'
        return page(active_tab=file_type, preview={'file_type': file_type, 'rows': rows[:20],
                    'columns': SCHEMAS[file_type], 'row_count': len(rows), 'token': token},
                    message=message)

    @app.post('/insert')
    def insert():
        check_csrf()
        if repo is None or not app.config['UI_IMPORTS_ENABLED']:
            abort(503, description='Import interactif non activé. Vérifiez la configuration et le contrôle d’accès. Aucun import effectué.')
        file_type = request.form.get('file_type')
        preview_id, data = store().load(request.form.get('temp_file', ''), browser_state()[0], file_type)
        rows = validate_rows(data['rows'], data['file_type'], max_rows=app.config['MAX_IMPORT_ROWS'])
        receipt = repo.import_rows(data['file_type'], rows, 'ui:' + preview_id)
        return page(active_tab=data['file_type'], receipt=receipt,
                    message='Import déjà confirmé.' if receipt['status'] == 'already_imported' else 'Import confirmé et enregistré.')

    @app.post('/cancel')
    def cancel():
        check_csrf()
        store().delete(request.form.get('temp_file', ''), browser_state()[0])
        return redirect(url_for('index'))

    @app.post('/api/v1/imports')
    def api_import():
        expected = app.config.get('IMPORT_API_TOKEN')
        if not expected or repo is None:
            abort(503, description='Import API is not configured.')
        actual = request.headers.get('Authorization', '')
        if not hmac.compare_digest(actual.encode(), ('Bearer ' + expected).encode()):
            abort(401, description='Valid bearer authentication required.')
        key = request.headers.get('Idempotency-Key', '')
        if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_.:/-]{0,195}', key):
            abort(400, description='Idempotency-Key is required (1–196 safe ASCII characters).')
        if not request.is_json:
            abort(415, description='Content-Type application/json required.')
        data = request.get_json()
        if not isinstance(data, dict) or set(data) != {'file_type', 'rows'}:
            raise ValidationError('Objet attendu : file_type et rows uniquement.')
        rows = validate_rows(data['rows'], data['file_type'], max_rows=app.config['MAX_IMPORT_ROWS'])
        receipt = {**repo.import_rows(data['file_type'], rows, 'api:' + key), 'rows_received': len(data['rows'])}
        return jsonify(receipt), 201 if receipt['status'] == 'imported' else 200

    @app.get('/download/template/<name>.<ext>')
    def download_template(name, ext):
        kind = {'edi_template': 'EDI', 'delivery_template': 'LIVRAISON'}.get(name)
        if not kind or ext not in {'csv', 'xlsx'}:
            abort(404)
        headers = SCHEMAS[kind]
        if ext == 'csv':
            data = ('\ufeff' + ','.join(headers) + '\r\n').encode('utf-8')
            mimetype = 'text/csv'
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = 'Template'
            ws.append(headers)
            ws.freeze_panes = 'A2'
            for cell in ws[1]:
                cell.comment = Comment('Conservez cet en-tête. Dates : AAAA-MM-JJ ; semaines EDI : AAAA-WSS. Les codes sont du texte.', 'EDI Stock')
                ws.column_dimensions[cell.column_letter].width = max(18, len(cell.value) + 3)
            for row in ws.iter_rows(min_row=2, max_row=201, max_col=len(headers)):
                for cell in row:
                    cell.number_format = '@'
            stream = io.BytesIO()
            wb.save(stream)
            data = stream.getvalue()
            mimetype = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        return Response(data, mimetype=mimetype, headers={'Content-Disposition': f'attachment; filename="{name}.{ext}"'})

    return app
